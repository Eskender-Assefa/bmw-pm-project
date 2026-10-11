"""Bounded, local PDF/spreadsheet extraction. Files and text are never instructions."""
import csv
import hashlib
import re
from pathlib import Path
from zipfile import ZipFile
import openpyxl
import pdfplumber
from sqlalchemy import select
from .audit import append_event
from .config import DATA_DIR
from .ingestion import evidence
from .models import Source
from .rag_schemas import DocumentMetadata

FORMATS = {'.pdf', '.xlsx', '.csv', '.txt', '.md'}
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50000
MAX_TEXT = 200000


def extract_rows(rows, sheet, metadata):
    header = None
    for number, row in enumerate(rows, 1):
        if number > MAX_ROWS:
            raise ValueError('A sheet exceeds 50,000 rows. Split the workbook before importing.')
        if len(row) > 200:
            raise ValueError('A sheet exceeds 200 columns.')
        values = [str(v) if v is not None else '' for v in row]
        if number < metadata.header_row:
            continue
        if number == metadata.header_row:
            header = values
            continue
        if not any(v.strip() for v in values):
            continue
        model = metadata.model
        for i, label in enumerate(header or []):
            if label.strip().lower() in ('model', 'vehicle_model', 'model_code') and i < len(values) and values[i].strip():
                row_model = values[i].strip().upper()
                if len(row_model) > 64 or not re.fullmatch(r'[A-Z0-9_ /-]+', row_model):
                    raise ValueError(f'Invalid vehicle model in {sheet}, row {number}.')
                if model and model != row_model:
                    raise ValueError(f'Model metadata disagrees with {sheet}, row {number}. Leave the file model blank for a multi-model sheet.')
                model = row_model
        body = '\n'.join(f'{(header[i] if header and i < len(header) and header[i] else "Column " + str(i+1))}: {value}'
                         for i, value in enumerate(values) if value.strip())
        if len(body) > MAX_TEXT:
            raise ValueError(f'Row {number} exceeds the text limit. Split its content.')
        yield body, {'sheet': sheet, 'row': number, 'header_row': metadata.header_row,
                     'range': f'A{number}:{openpyxl.utils.get_column_letter(len(values))}{number}'}, model


def extract_document(path, metadata, warnings):
    extension = path.suffix.lower()
    if extension == '.pdf':
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > 500:
                raise ValueError('PDF exceeds 500 pages. Split the document.')
            for number, page in enumerate(pdf.pages, 1):
                body = page.extract_text() or ''
                if len(body) > MAX_TEXT:
                    raise ValueError(f'PDF page {number} exceeds the text limit.')
                if not body.strip():
                    warnings.append(f'Page {number}: no selectable text; OCR may be needed.')
                else:
                    yield body, {'page': number, 'extraction_method': 'pdfplumber-text-v1'}, metadata.model
    elif extension == '.xlsx':
        with ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Expanded workbook exceeds 100 MB.')
        book = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        warnings.append('Excel uses cached formula values; formulas are not recalculated. First configured header row is retained with each row.')
        try:
            for sheet in book.worksheets:
                yield from extract_rows(sheet.iter_rows(values_only=True), sheet.title, metadata)
        finally:
            book.close()
    elif extension == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as handle:
            yield from extract_rows(csv.reader(handle), 'CSV', metadata)
    else:
        with path.open(encoding='utf-8-sig') as handle:
            for number, line in enumerate(handle, 1):
                if number > MAX_ROWS or len(line) > MAX_TEXT:
                    raise ValueError('Text file exceeds line or text limits.')
                if line.strip():
                    yield line.rstrip('\r\n'), {'line': number}, metadata.model


def import_documents(session, payload, progress=lambda _: None):
    metadata = DocumentMetadata.model_validate(payload['metadata'])
    if metadata.model:
        metadata.model = metadata.model.strip().upper() or None
    imported, skipped, failed = [], [], []
    root = (DATA_DIR / 'uploads').resolve()
    for item in payload['files']:
        filename = item['filename']
        try:
            path = Path(item['path']).resolve()
            if path.parent != root or path.suffix.lower() not in FORMATS or not path.is_file() or path.stat().st_size > MAX_BYTES:
                raise ValueError('Invalid staged upload or file size.')
            checksum = hashlib.sha256(path.read_bytes()).hexdigest()
            prior = session.scalar(select(Source).where(Source.checksum == checksum))
            if prior:
                skipped.append({'source_id': prior.id, 'filename': filename, 'reason': 'Identical bytes already imported; existing metadata and reviews were preserved.'})
                continue
            progress('Extracting ' + filename)
            warnings = []
            # Per-file savepoints allow a bad file to fail without losing other files.
            with session.begin_nested():
                source = Source(id=checksum[:32], checksum=checksum, filename=filename,
                                format=path.suffix.lower()[1:], kind='document', origin=metadata.origin,
                                language='unknown', market=metadata.market, path=str(path),
                                extraction_version='document-v1', record_count=0, row_count=0)
                session.add(source)
                session.flush()
                count = 0
                for body, locator, model in extract_document(path, metadata, warnings):
                    count += 1
                    if count > MAX_ROWS:
                        raise ValueError('File exceeds 50,000 extracted records.')
                    evidence(session, source, locator, body, 'document', model)
                if not count:
                    raise ValueError('No extractable records. A scanned PDF needs OCR; a spreadsheet needs headers and data rows.')
                source.record_count = source.row_count = count
                append_event(session, 'Demo PM', 'document.imported', source.id,
                             {'checksum': checksum, 'records': count, 'metadata': metadata.model_dump(), 'warnings': warnings})
            imported.append({'source_id': source.id, 'filename': filename, 'records': count, 'warnings': warnings})
        except Exception as exc:
            failed.append({'filename': filename, 'error': str(exc)[:500]})
    from .rag import index_sources
    indexed = index_sources(session, [i['source_id'] for i in imported + skipped], progress) if imported or skipped else {'added': 0}
    return {'imported': imported, 'skipped': skipped, 'failed': failed, 'index': indexed}
