"""Deterministic source adapters. Imports never use AI or execute document text."""
import hashlib
import io
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
import pdfplumber
from pptx import Presentation
from openpyxl import load_workbook
from sqlalchemy import select, text
from .config import DATA_DIR, PIPELINE_VERSION, MAPPING_VERSION
from .models import Source, EvidenceRef, Feedback, Aspect, Fact, Vehicle, EvaluationLabel
from .taxonomy import seed_taxonomy, feature_for_tag, serious_concern
from .audit import append_event, digest

EXPECTED = {'F70': 2128, 'G60': 4000, 'G70': 4000}
ALLOWED_FILES = {f'{m}_feedback_hackathon.xlsx' for m in EXPECTED} | {f'{m}_OptionList.PDF' for m in EXPECTED} | {'sales_volumes.xlsx','F70_G60_G68_G70_customer_studies.xlsx'}

def evidence(session, source, locator, content, kind, model=None):
    item = EvidenceRef(id=digest([source.id, locator])[:32], source_id=source.id, locator=locator, text=content,
                       kind=kind, origin=source.origin, model=model)
    session.add(item)
    session.flush()
    session.execute(text('INSERT INTO evidence_fts(evidence_id, text) VALUES (:id, :text)'), {'id': item.id, 'text': content})
    return item

def import_archive(session, archive, progress=lambda _: None):
    seed_taxonomy(session)
    summary = {'imported': [], 'skipped': []}
    with zipfile.ZipFile(archive) as z:
        members = [n for n in z.namelist() if not n.endswith('/')]
        names = [Path(n).name for n in members]
        if len(members) != 8 or set(names) != ALLOWED_FILES or len(set(names)) != 8:
            raise ValueError('Expected the eight supplied BMW source files. The archive inventory does not match.')
        for member in members:
            info = z.getinfo(member)
            if info.file_size > 30_000_000:
                raise ValueError('Source exceeds the 30 MB import limit.')
            name = Path(member).name
            progress(f'Importing {name}')
            data = z.read(member)
            checksum = hashlib.sha256(data).hexdigest()
            if session.scalar(select(Source).where(Source.checksum == checksum)):
                summary['skipped'].append(name)
                continue
            if session.scalar(select(Source).where(Source.filename == name, Source.active == True)):
                raise ValueError(f'{name} has a new checksum. Source-edition replacement requires an explicit migration; existing decisions are preserved.')
            path = DATA_DIR / 'sources' / f'{checksum[:16]}-{name}'
            path.write_bytes(data)
            kind = 'customer_statement' if 'feedback_hackathon' in name else 'published_specification' if name.lower().endswith('.pdf') else 'aggregate_study' if 'studies' in name else 'sales_aggregate'
            source = Source(id=checksum[:32], filename=name, checksum=checksum, format=Path(name).suffix[1:].lower(),
                            kind=kind, origin='unknown', language='de' if kind=='published_specification' else 'en',
                            market='DE' if kind=='published_specification' else None, path=str(path), extraction_version=PIPELINE_VERSION)
            session.add(source)
            session.flush()
            if kind == 'customer_statement':
                import_feedback(session, source, data)
            elif kind == 'published_specification':
                import_pdf(session, source, data)
            elif kind == 'aggregate_study':
                import_studies(session, source, data)
            else:
                import_sales(session, source, data)
            summary['imported'].append({'filename': name, 'records': source.record_count, 'rows': source.row_count})
        create_evaluation_set(session)
        append_event(session, 'Importer', 'sources.imported', 'archive', summary)
    return summary

def import_files(session, paths, progress=lambda _: None):
    """Import a supplied subset and retain reference documents as searchable text."""
    reference_files = {'BMW_G60_Scopes_and_Data.pdf', '261010_BMW_Product_Management_Key_Note(1).pptx'}
    paths = [Path(path) for path in paths]
    for path in paths:
        if path.name not in ALLOWED_FILES | reference_files:
            raise ValueError(f'Unsupported source file: {path.name}')
        if not path.is_file() or path.stat().st_size > 30_000_000:
            raise ValueError(f'Source must be an existing file under 30 MB: {path.name}')
    seed_taxonomy(session)
    summary = {'imported': [], 'skipped': []}
    for path in paths:
        name = path.name
        progress(f'Importing {name}')
        data = path.read_bytes()
        checksum = hashlib.sha256(data).hexdigest()
        if session.scalar(select(Source).where(Source.checksum == checksum)):
            summary['skipped'].append(name)
            continue
        if session.scalar(select(Source).where(Source.filename == name, Source.active == True)):
            raise ValueError(f'{name} has a new checksum. Source-edition replacement requires an explicit migration; existing decisions are preserved.')
        reference = name in reference_files
        kind = ('project_reference' if reference else 'customer_statement' if 'feedback_hackathon' in name
                else 'published_specification' if path.suffix.lower() == '.pdf'
                else 'aggregate_study' if 'studies' in name else 'sales_aggregate')
        retained = DATA_DIR / 'sources' / f'{checksum[:16]}-{name}'
        retained.write_bytes(data)
        source = Source(id=checksum[:32], filename=name, checksum=checksum, format=path.suffix[1:].lower(),
                        kind=kind, origin='unknown', language='de' if kind == 'published_specification' else 'en',
                        market='DE' if kind == 'published_specification' else None,
                        path=str(retained), extraction_version=PIPELINE_VERSION)
        session.add(source)
        session.flush()
        if reference:
            if source.format == 'pdf':
                with pdfplumber.open(io.BytesIO(data)) as pdf:
                    for number, page in enumerate(pdf.pages, 1):
                        evidence(session, source, {'page': number, 'bbox': [0, 0, page.width, page.height]},
                                 page.extract_text(x_tolerance=2) or '', kind, 'G60')
                    source.record_count = source.row_count = len(pdf.pages)
            else:
                deck = Presentation(io.BytesIO(data))
                for number, slide in enumerate(deck.slides, 1):
                    blocks = []
                    def collect(shapes):
                        for shape in shapes:
                            if shape.has_text_frame:
                                blocks.append(shape.text)
                            if shape.has_table:
                                blocks.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
                            if hasattr(shape, 'shapes'):
                                collect(shape.shapes)
                    collect(slide.shapes)
                    evidence(session, source, {'slide': number}, '\n'.join(blocks), kind, 'G60')
                source.record_count = source.row_count = len(deck.slides)
        elif kind == 'customer_statement':
            import_feedback(session, source, data)
        elif kind == 'published_specification':
            import_pdf(session, source, data)
        elif kind == 'aggregate_study':
            import_studies(session, source, data)
        else:
            import_sales(session, source, data)
        summary['imported'].append({'filename': name, 'records': source.record_count, 'rows': source.row_count})
    create_evaluation_set(session)
    append_event(session, 'Importer', 'sources.imported', 'files', summary)
    return summary


def import_feedback(session, source, data):
    w = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = w['Feedback_Explorer']
    records = defaultdict(list)
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if row[0]:
            records[str(row[0])].append((index, row))
    model = source.filename[:3]
    if len(records) != EXPECTED[model]:
        raise ValueError(f'{model}: expected {EXPECTED[model]} IDs, found {len(records)}.')
    hashes = Counter(hashlib.sha256(str(rows[0][1][8]).encode()).hexdigest() for rows in records.values())
    for external_id, rows in records.items():
        first = rows[0][1]
        if len({tuple(r[:7]) for _, r in rows}) != 1 or len({r[8] for _,r in rows}) != 1:
            raise ValueError(f'Conflicting identity/text for {external_id}; review before importing.')
        raw_text = str(first[8])
        text_hash = hashlib.sha256(raw_text.encode()).hexdigest()
        record = Feedback(id=f'{source.id[:12]}:{external_id}', source_id=source.id, external_id=external_id, model=model,
                          country=str(first[4]), powertrain={'BEVE':'BEV'}.get(first[6], str(first[6])), raw_powertrain=str(first[6]),
                          text=raw_text, text_hash=text_hash, source_label=str(first[1]), tags=[], duplicate_text=hashes[text_hash]>1)
        session.add(record)
        session.flush()
        tags, seen = [], set()
        for row_number, row in rows:
            ref = evidence(session, source, {'sheet':sheet.title, 'row':row_number, 'range':f'A{row_number}:L{row_number}', 'external_id':external_id}, raw_text, 'customer_statement', model)
            kind = str(row[7] or 'Unclassified')
            tags.append({'row':row_number,'type':kind,'level2':row[10],'level3':row[11],'combined':row[9], 'evidence_id':ref.id})
            feature = feature_for_tag(session, row[10])
            key = (feature.id, kind)
            if key not in seen:
                seen.add(key)
                session.add(Aspect(feedback_id=record.id, evidence_id=ref.id, feature_id=feature.id, feedback_type=kind,
                                   quote=raw_text, mapping_status='pending' if 'other_pending_review' in feature.mapping_reason else 'suggested',
                                   mapping_reason=feature.mapping_reason, mapping_version=MAPPING_VERSION,
                                   critical=serious_concern(raw_text, kind) and (feature.category in ('C04','C09') or feature.id=='drivetrain')))
        record.tags = tags
    source.record_count = len(records)
    source.row_count = sum(len(r) for r in records.values())
    session.flush()
    w.close()

PRICE_GROUPS = {
 'F70': [(['116','120','123 xDrive','M135 xDrive'],'ICE'), (['118d','120d'],'ICE')],
 'G60': [(['i5 eDrive40','i5 xDrive40','i5 M60 xDrive'],'BEV'), (['530e','530e xDrive','550e xDrive'],'PHEV'), (['520i','520d','520d xDrive','540d xDrive'],'ICE')],
 'G70': [(['i7 50xDrive','i7 60xDrive','i7 M70 xDrive'],'BEV'), (['750e xDrive','M760e xDrive'],'PHEV'), (['740 xDrive','740d xDrive'],'ICE')],
}

def number_value(value):
    cleaned = value.replace(' ', '').replace('\u00a0','')
    match = re.search(r'\d[\d.]*,?\d*', cleaned)
    if not match:
        return None
    return float(match.group().replace('.', '').replace(',', '.'))

def vehicle(session, model, variant, powertrain=None):
    key = f'{model}:DE:{variant}'
    if not session.get(Vehicle, key):
        session.add(Vehicle(id=key, model=model, name=variant, variant=variant, powertrain=powertrain, market='DE', aliases=[variant]))
        session.flush()
    return key

def import_pdf(session, source, data):
    model = source.filename[:3]
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        source.edition = 'July 2026, issue 2' if model=='G60' else 'July 2026, issue 1'
        source.source_date = '2026-06-22' if model=='G60' else '2026-05-28'
        for page_number, page in enumerate(pdf.pages, 1):
            content = page.extract_text(x_tolerance=2) or ''
            page_ref = evidence(session, source, {'page':page_number, 'bbox':[0,0,page.width,page.height]}, content, 'published_specification', model)
            if page_number in (4,5):
                extract_prices(session, source, model, page, page_number)
            extract_options(session, source, model, page, page_number)
        source.row_count = len(pdf.pages)
        session.flush()
        source.record_count = len(session.scalars(select(Fact).join(EvidenceRef, Fact.evidence_id==EvidenceRef.id).where(EvidenceRef.source_id==source.id)).all())

def extract_prices(session, source, model, page, page_number):
    words = [w for w in page.extract_words(x_tolerance=2) if w['upright']]
    for names, powertrain in PRICE_GROUPS[model]:
        anchors = []
        for name in names:
            tokens = name.split()
            matches = []
            for i in range(len(words)-len(tokens)+1):
                group = words[i:i+len(tokens)]
                if [re.sub(r'[¹²³⁴⁵]+','',w['text']) for w in group] == tokens and group[0]['x0']>180 and max(w['top'] for w in group)-min(w['top'] for w in group)<3:
                    matches.append((sum((w['x0']+w['x1'])/2 for w in group)/len(group), group[0]['top']))
            if matches:
                anchors.append((name, *matches[0]))
        if len(anchors)!=len(names) or max(a[2] for a in anchors)-min(a[2] for a in anchors)>5:
            continue
        anchors.sort(key=lambda a:a[1])
        header_y = max(a[2] for a in anchors)
        price_labels = [w for w in words if w['text'].startswith('FAHRZEUGGRUNDPREIS') and header_y < w['top'] < header_y+250]
        if not price_labels:
            continue
        end_y = min(w['top'] for w in price_labels)+24
        for label_pattern, label, category, unit in [
            ('FAHRZEUGGRUNDPREIS','Base price','C03','EUR'),
            ('Leistung','System power','C01','kW'),
            ('Energieverbrauch,','Electricity consumption (WLTP)','C02','kWh/100 km'),
            ('Kraftstoffverbrauch,','Fuel consumption (WLTP)','C02','L/100 km'),
            ('Getriebe','Transmission','C01',None)]:
            rows = [w for w in words if w['text'].startswith(label_pattern) and w['x0']<100 and header_y<w['top']<end_y]
            if not rows: continue
            target_y = rows[0]['top']
            for i,(name,center,_) in enumerate(anchors):
                left=(anchors[i-1][1]+center)/2 if i else 195
                right=(anchors[i+1][1]+center)/2 if i<len(anchors)-1 else page.width-35
                cell_words = [w for w in words if left<w['x0']<right and abs(w['top']-target_y)<3]
                raw = ' '.join(w['text'] for w in cell_words)
                if not raw or raw.strip() in ('–','-'): continue
                value = raw if unit is None else number_value(raw)
                if value is None: continue
                bbox = [left,target_y-1,right,target_y+12]
                excerpt = f'{name}: {label} = {raw}'
                ref = evidence(session, source, {'page':page_number,'bbox':bbox,'column':name,'row_label':label}, excerpt, 'published_specification', model)
                conditions={'raw_value':raw,'source_edition':source.edition, 'market':'DE'}
                if unit=='EUR': conditions.update({'tax_basis':'includes 19% VAT','delivery_costs':'excluded'})
                if 'WLTP' in label:
                    conditions['test_method']='WLTP'
                    numbers = re.findall(r'\d+(?:,\d+)?', raw)
                    if len(numbers)>1 and ('–' in raw or '-' in raw): value={'min':min(float(n.replace(',','.')) for n in numbers),'max':max(float(n.replace(',','.')) for n in numbers)}
                session.add(Fact(evidence_id=ref.id, vehicle_id=vehicle(session,model,name,powertrain), model=model, category=category,
                                 label=label, value=value, unit=unit, market='DE', variant=name, conditions=conditions,
                                 review_status='pending', evidence_kind='published_specification', origin='unknown'))

def extract_options(session, source, model, page, page_number):
    words = page.extract_words(x_tolerance=2)
    rotated = defaultdict(list)
    for w in words:
        if not w['upright'] and w['top']<150:
            rotated[round(w['x0'])].append(w)
    headers = [(x,' '.join(w['text'][::-1] for w in sorted(ws,key=lambda w:w['top'],reverse=True))) for x,ws in rotated.items()]
    code_cols=[x for x,t in headers if 'SA-Nr' in t]
    variants=[(x,t) for x,t in headers if re.match(r'^(?:i[57] |\d{3}[a-z]?|M\d{3})',t)]
    if not code_cols or not variants: return
    code_x=code_cols[0]
    code_words=[w for w in words if w['upright'] and abs(w['x0']-code_x)<12 and re.fullmatch(r'[A-Z0-9]{3}',w['text']) and any(c.isdigit() for c in w['text']) and 130<w['top']<770]
    code_words.sort(key=lambda w:w['top'])
    for i,code in enumerate(code_words):
        top=code['top']; bottom=code_words[i+1]['top']-2 if i+1<len(code_words) else min(top+120,770)
        description_words=[w for w in words if w['upright'] and w['x0']<code_x-12 and top-2<=w['top']<bottom]
        description=' '.join(w['text'] for w in description_words)
        title=' '.join(w['text'] for w in description_words if abs(w['top']-top)<4)
        if not title: continue
        prices=[w['text'] for w in words if w['upright'] and code_x+15<w['x0']<min(v[0] for v in variants)-8 and abs(w['top']-top)<4]
        raw_price=' '.join(prices)
        category='C09' if any(t in title.lower() for t in ('assistant','assistent','driving')) else 'C05' if any(t in title.lower() for t in ('sitz','komfort','klima','sound','heiz')) else 'C07' if any(t in title.lower() for t in ('navigation','connected','display','telefon','wlan')) else 'C04' if any(t in title.lower() for t in ('fahrwerk','bremse','lenkung')) else 'C10'
        for x,name in variants:
            marks=[w['text'] for w in words if w['upright'] and abs(w['x0']-x)<10 and abs(w['top']-top)<4]
            availability='standard' if '■' in marks else 'optional' if any(m in ('□','▫','◻') for m in marks) else 'unknown'
            bbox=[40,top-2,page.width-35,bottom]
            ref=evidence(session,source,{'page':page_number,'bbox':bbox,'option_code':code['text'],'column':name},f'{title}\n{description}\nOption code: {code["text"]}; printed price: {raw_price}; symbols: {" ".join(marks)}','published_specification',model)
            session.add(Fact(evidence_id=ref.id,vehicle_id=vehicle(session,model,name),model=model,category=category,label=title,
                             value=availability,market='DE',variant=name,knowledge_state='unknown' if availability=='unknown' else 'known',
                             conditions={'option_code':code['text'],'price_eur_including_vat':number_value(raw_price),'raw_price':raw_price,'tax_basis':'19% VAT','conditions_text':description,'availability_is_candidate':True},
                             evidence_kind='published_specification',review_status='pending',origin='unknown'))

def import_sales(session, source, data):
    w=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    count=0
    for sheet in w:
        for i,row in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
            if not row[0]: continue
            for col,year in [(2,2024),(3,2025),(4,2030)]:
                ref=evidence(session,source,{'sheet':sheet.title,'row':i,'range':f'{chr(65+col)}{i}'},f'{sheet.title}, {row[0]}, volume_{year}: {row[col]}','sales_aggregate',sheet.title)
                session.add(Fact(evidence_id=ref.id,vehicle_id=sheet.title,model=sheet.title,category='C10',label=f'Sales volume {year}',value=row[col],unit='vehicles',market=row[1],conditions={'year':year,'aggregation':'G60 and G68 combined' if sheet.title=='G60_G68' else 'model total','period_kind':'future value supplied; basis unverified' if year==2030 else 'historical value supplied; basis unverified'},review_status='source_transcribed',evidence_kind='sales_aggregate'))
                count+=1
    source.record_count=count; source.row_count=12
    w.close()

def import_studies(session, source, data):
    w=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    count=0
    sheet=w['CN_EU_2025']
    columns={2:('G68','CN'),3:('G70','CN'),4:('F70','EU'),5:('G60','EU'),6:('G70','EU')}
    label=''
    for i,row in enumerate(sheet.iter_rows(values_only=True),1):
        if i<3: continue
        if row[0] and row[0]!='Mean': label=row[0]
        if row[0]!='Mean': continue
        for col,(model,market) in columns.items():
            if not isinstance(row[col-1],(float,int)): continue
            ref=evidence(session,source,{'sheet':sheet.title,'row':i,'range':f'{chr(64+col)}{i}'},f'{model}, {market}, {label}, Mean: {row[col-1]}','study_mean',model)
            feature=feature_for_tag(session,label)
            session.add(Fact(evidence_id=ref.id,vehicle_id=model,model=model,category=feature.category,label=label,value=row[col-1],unit='source mean',market=market,conditions={'year':2025,'scale':None,'sample_size':None,'instrument':'CN_EU_2025'},review_status='source_transcribed',evidence_kind='study_mean'))
            count+=1
    sheet=w['US_2025']; label=''
    for i,row in enumerate(sheet.iter_rows(values_only=True),1):
        if row[0] and row[1] is None: label=str(row[0]); continue
        if not row[0] or row[0]=='Sample total': continue
        for col,model in [(3,'G60'),(4,'G70')]:
            if not isinstance(row[col-1],(float,int)): continue
            ref=evidence(session,source,{'sheet':sheet.title,'row':i,'range':f'{chr(64+col)}{i}'},f'{model}, US, {label}, {row[0]}: {row[col-1]}','study_distribution',model)
            feature=feature_for_tag(session,label)
            session.add(Fact(evidence_id=ref.id,vehicle_id=model,model=model,category=feature.category,label=label,value=row[col-1],unit='proportion',market='US',conditions={'year':2025,'response':row[0],'sample_size':None,'instrument':'US_2025'},review_status='source_transcribed',evidence_kind='study_distribution'))
            count+=1
    source.record_count=count; source.row_count=count
    w.close()

def create_evaluation_set(session):
    if session.scalar(select(EvaluationLabel.feedback_id).limit(1)): return
    # Balanced round-robin across model/country/source type. Identical text stays in one split.
    records=session.scalars(select(Feedback).order_by(Feedback.id)).all()
    groups=defaultdict(list)
    for record in records:
        kinds=sorted({t['type'] for t in record.tags})
        groups[(record.model,record.country,','.join(kinds))].append(record)
    picked=[]; hashes=set()
    while len(picked)<100 and any(groups.values()):
        for key in sorted(groups):
            if not groups[key]: continue
            record=groups[key].pop(0)
            if record.text_hash in hashes: continue
            hashes.add(record.text_hash); picked.append(record)
            if len(picked)==100: break
    for i,record in enumerate(picked):
        session.add(EvaluationLabel(feedback_id=record.id,split='calibration' if i<40 else 'held_out'))
