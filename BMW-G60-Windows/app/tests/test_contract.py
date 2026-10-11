import json
import os
import sqlite3
import socket
import subprocess
import sys
from pathlib import Path
from copy import deepcopy
import pytest
import httpx
from sqlalchemy import select,func,text
from sqlalchemy.exc import IntegrityError
from fastapi import HTTPException
from pptx import Presentation
from backend.models import Source,Feedback,Aspect,Feature,Fact,EvidenceRef,RequirementRevision,Decision,AuditEvent,Job,Export,EvaluationLabel,AnalysisRun
from backend.schemas import Cohort,DecisionRequest,AnalysisRequest,ExportRequest
from backend.services import ranking,decide,detail,evaluation_metrics,seed_requirements
from backend.analysis import prepare_input,validate_output,analyze
from backend.audit import verify_audit,digest
from backend.db import SessionLocal
from backend.config import ROOT,DATA_DIR
from backend.connectors import import_web, normalize_url
from backend.schemas import WebConnectorRequest

def target(session):
    return next(r for r in ranking(session,Cohort()).requirements if r.feature=='Touch screen, operation')

def decision(req,action,**kwargs):
    return DecisionRequest(cohort=Cohort(),expected_revision=req.revision,action=action,actor='Automated contract test',**kwargs)

def valid_proposal(inputs):
    ref=inputs['evidence'][0]
    return {'title':'Make touchscreen controls easier to use','description':'Provide an understandable interaction for the cited touchscreen task.','acceptance_criteria':['Proposed: validate the interaction with the affected customers.'],'evidence_ids':[ref['id']],'uncertainties':['Variant is unverified.'],
      'aspects':[{'evidence_id':ref['id'],'quote':ref['text'],'category':inputs['category'],'feature_id':inputs['feature_id'],'feedback_type':'Difficult to Use','sentiment':-30,'explanation':'Interpretation for a test fixture, not a real AI generation.'}]}

def test_import_counts_and_original_rows(session):
    assert dict(session.execute(select(Feedback.model,func.count()).group_by(Feedback.model)).all())=={'F70':2128,'G60':4000,'G70':4000}
    sources=session.scalars(select(Source)).all()
    assert len(sources)==8
    assert sum(s.row_count for s in sources if s.kind=='customer_statement')==13771
    assert session.scalar(select(func.count()).select_from(EvidenceRef).where(EvidenceRef.kind=='customer_statement'))==13771
    assert all(Path(s.path).exists() and len(s.checksum)==64 for s in sources)
    assert session.scalar(select(func.count()).select_from(Feedback).where(Feedback.source_date!=None))==0

def test_unchanged_reimport_is_idempotent(session):
    from backend.ingestion import import_archive
    archive=Path(os.getenv('BMWPM_SOURCE_ARCHIVE',r'C:\Users\rider\Downloads\BMW-Data-20261010T183823Z-1-001.zip'))
    if not archive.exists():pytest.skip('Source archive not available; set BMWPM_SOURCE_ARCHIVE.')
    before=[session.scalar(select(func.count()).select_from(m)) for m in (Source,Feedback,Aspect,EvidenceRef)]
    import_archive(session,archive,lambda _:None);session.commit()
    assert before==[session.scalar(select(func.count()).select_from(m)) for m in (Source,Feedback,Aspect,EvidenceRef)]

def test_denominator_is_cohort_correct_and_independent_of_category(session):
    all_rows=ranking(session,Cohort()); software=ranking(session,Cohort(),category='C07')
    assert all_rows.denominator==software.denominator==3610
    full={r.id:r for r in all_rows.requirements}
    assert all(r.numerator==len(set(r.supporting_ids)) and r.denominator==3610 and r.demand==r.numerator/3610 for r in software.requirements)
    assert all(r.numerator==full[r.id].numerator for r in software.requirements)
    bev=ranking(session,Cohort(powertrain='BEV'))
    assert bev.denominator==session.scalar(select(func.count()).select_from(Feedback).where(Feedback.model=='G60',Feedback.country=='US',Feedback.powertrain=='BEV'))
    assert ranking(session,Cohort(origin='synthetic')).denominator==0
    assert ranking(session,Cohort(model='G68')).denominator==0

def test_mixed_comments_and_identical_text_ids(session):
    mixed=session.execute(select(Aspect.feedback_id,func.count(func.distinct(Aspect.feature_id))).group_by(Aspect.feedback_id).having(func.count(func.distinct(Aspect.feature_id))>1)).first()
    assert mixed is not None
    duplicates=session.scalars(select(Feedback).where(Feedback.duplicate_text==True)).all()
    assert duplicates and len({f.id for f in duplicates})==len(duplicates)
    assert len({f.text_hash for f in duplicates})<len(duplicates)
    for req in ranking(session,Cohort()).requirements:
        assert req.numerator==len(set(req.supporting_ids))
        for sample in req.evidence_samples:
            ref=session.get(EvidenceRef,sample['id'])
            assert sample['quote'] in ref.text and ref.locator['row']>=2

def test_studies_sales_and_catalog_retain_boundaries(session):
    means=session.scalars(select(Fact).where(Fact.evidence_kind=='study_mean')).all()
    distributions=session.scalars(select(Fact).where(Fact.evidence_kind=='study_distribution')).all()
    assert means and distributions
    assert all(f.conditions.get('sample_size') is None for f in means+distributions)
    combined=session.scalars(select(Fact).where(Fact.model=='G60_G68')).all()
    assert combined and all(f.evidence_kind=='sales_aggregate' for f in combined)
    assert not session.scalar(select(Fact).where(Fact.model=='G60',Fact.evidence_kind=='sales_aggregate'))
    power=session.scalar(select(Fact).where(Fact.model=='G60',Fact.variant=='i5 eDrive40',Fact.label=='System power'))
    price=session.scalar(select(Fact).where(Fact.model=='G60',Fact.variant=='i5 eDrive40',Fact.label=='Base price'))
    assert power.value==250 and power.unit=='kW' and price.value==70900
    assert power.market==price.market=='DE' and power.review_status=='pending'
    ref=session.get(EvidenceRef,power.evidence_id)
    assert ref.locator['page']==4

def test_critical_sort_precedes_reviewed_bands(session):
    rows=ranking(session,Cohort()).requirements
    critical=[r for r in rows if r.critical]
    assert critical and rows[:len(critical)]==critical
    assert all(r.band=='Unscored' for r in rows)

def test_decision_revision_persistence_and_cohort_isolation(session):
    req=target(session);session.rollback()
    decide(session,req.id,decision(req,'approve',band='High',rationale='Review evidence and validate the affected configuration.'))
    with SessionLocal() as reopened:
        saved,_=detail(reopened,req.id,Cohort()); assert saved.status=='approved' and saved.band=='High'
        # A band for all engines cannot silently become a BEV decision.
        bev,_=detail(reopened,req.id,Cohort(powertrain='BEV')); assert bev.status=='proposed' and bev.band=='Unscored'
    session.rollback()
    edited=decide(session,req.id,decision(req,'edit',title='Make common touchscreen actions easy to find',rationale='Clarify desired customer behavior and preserve evidence.'))
    assert edited['revision']==2 and edited['status']=='proposed' and edited['band']=='Unscored'
    assert session.scalar(select(func.count()).select_from(RequirementRevision).where(RequirementRevision.requirement_id==req.id))==2
    assert session.scalar(select(func.count()).select_from(Decision).where(Decision.requirement_id==req.id))==2
    assert verify_audit(session)['valid']
    session.rollback()
    with pytest.raises(HTTPException) as exc:decide(session,req.id,decision(req,'approve'))
    assert exc.value.status_code==409

@pytest.mark.parametrize('action',['reject','postpone','set_band','edit'])
def test_missing_rationale_is_atomic(session,action):
    req=target(session);before=session.scalar(select(func.count()).select_from(Decision));session.rollback()
    with pytest.raises(HTTPException):decide(session,req.id,decision(req,action))
    session.rollback()
    assert session.scalar(select(func.count()).select_from(Decision))==before
    assert detail(session,req.id,Cohort())[0].revision==req.revision

def test_append_only_guards_and_tamper_detection(session):
    assert verify_audit(session)['valid']
    with pytest.raises(IntegrityError):session.execute(text("UPDATE audit_events SET actor='altered' WHERE seq=1"))
    session.rollback()
    # Simulate privileged tampering in this disposable test database.
    session.execute(text('DROP TRIGGER immutable_audit_events_update'))
    session.execute(text("UPDATE audit_events SET actor='altered' WHERE seq=1"));session.commit()
    assert verify_audit(session)['valid'] is False

def test_invalid_ai_references_quotes_and_feature_fail_closed(session):
    req=target(session);inputs=prepare_input(session,AnalysisRequest(cohort=Cohort(),requirement_id=req.id))
    output=valid_proposal(inputs);assert validate_output(inputs,output)
    for mutation in ('id','quote','feature'):
        bad=deepcopy(output)
        if mutation=='id':bad['evidence_ids']=['invented-id']
        elif mutation=='quote':bad['aspects'][0]['quote']='This quote was not supplied.'
        else:bad['aspects'][0]['feature_id']='invented-feature'
        with pytest.raises(ValueError):validate_output(inputs,bad)

def test_replay_missing_stale_and_live_failure_are_visible(session,monkeypatch):
    req=target(session);request=AnalysisRequest(cohort=Cohort(),requirement_id=req.id)
    # Different configured model cannot reuse an existing snapshot.
    monkeypatch.setenv('OPENAI_MODEL','test-unavailable-model')
    with pytest.raises(ValueError,match='No matching saved replay'):analyze(session,request)
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    request.mode='live'
    with pytest.raises(ValueError,match='not configured'):analyze(session,request)
    assert detail(session,req.id,Cohort())[0].revision==req.revision

def test_real_replay_provenance_and_stale_revision(session,monkeypatch):
    monkeypatch.setenv('OPENAI_MODEL','gpt-4.1-mini-2025-04-14')
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    req=target(session);request=AnalysisRequest(cohort=Cohort(),requirement_id=req.id)
    # Force portable-fixture import, without modifying the original database.
    session.execute(text('DROP TRIGGER immutable_analysis_runs_delete'))
    session.execute(text('DELETE FROM analysis_runs'));session.commit()
    result=analyze(session,request);session.commit()
    assert result['mode']=='replay' and result['original_mode']=='live'
    assert result['model']=='gpt-4.1-mini-2025-04-14' and result['sample_count']==8
    assert result['generated_at'].startswith('2026-10-10')
    session.rollback();decide(session,req.id,decision(req,'edit',title='Changed requirement makes the saved replay stale',rationale='Verify old input versions cannot be replayed.'))
    with pytest.raises(ValueError,match='No matching saved replay'):analyze(session,request)

def test_interrupted_jobs_require_explicit_retry(session):
    from backend.jobs import recover_jobs,execute
    job=Job(kind='unknown-test-kind',input={},state='running');session.add(job);session.commit();job_id=job.id
    recover_jobs();session.expire_all()
    assert session.get(Job,job_id).state=='interrupted'
    session.rollback();execute(job_id);session.expire_all()
    assert session.get(Job,job_id).state=='failed' and session.get(Job,job_id).error=='Unknown job type.'

def test_human_evaluation_separate_and_uncalibrated(session):
    labels=session.scalars(select(EvaluationLabel)).all()
    assert len(labels)==100
    assert sum(l.split=='held_out' for l in labels)==60
    assert all(l.expected_categories is None for l in labels)
    assert evaluation_metrics(session)['automatic_promotion'] is False
    assert evaluation_metrics(session)['splits']['held_out']['exact_category_accuracy'] is None

def test_api_search_contract_and_cross_origin_protection(client):
    response=client.get('/api/v1/evidence',params={'q':'touchscreen','model':'G60'})
    assert response.status_code==200 and response.json()['total']>0
    item=response.json()['items'][0]
    ref=client.get('/api/v1/evidence/'+item['id']).json()
    assert ref['source']['checksum'] and ref['locator']
    assert client.get('/api/v1/requirements?model=G60&country=US&category=C07').json()['denominator']==3610
    assert client.post('/api/v1/exports',json={'cohort':Cohort().model_dump()},headers={'origin':'https://untrusted.example'}).status_code==403
    assert client.get('/api/v1/evidence?q=%22%20OR%20%2A').status_code==200

def web_client_factory(pages):
    def handler(request):
        if request.url.path == '/robots.txt':
            body = pages.get('robots', 'User-agent: *\nAllow: /\n')
            return httpx.Response(200, text=body)
        page = pages.get(str(request.url))
        if isinstance(page, Exception):
            raise page
        if page is None:
            return httpx.Response(404, text='missing')
        body, content_type = page
        return httpx.Response(200, content=body.encode(), headers={'content-type': content_type, 'last-modified': 'Sat, 10 Oct 2026 12:00:00 GMT'})
    return httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=True)

def test_seeded_web_connector_imports_observed_review_evidence(session,monkeypatch):
    monkeypatch.setattr('backend.connectors.socket.getaddrinfo',lambda host,port:[(socket.AF_INET,socket.SOCK_STREAM,6,'',( '93.184.216.34',0))])
    html='<html><body><article><p>The BMW iDrive touchscreen is difficult to use while driving. I want clearer climate shortcuts because the screen menus take too much attention on every trip.</p><blockquote>I love the display quality, but the touch screen operation needs fewer layers for common controls.</blockquote></article></body></html>'
    request=WebConnectorRequest(urls=['https://owners.example/forum/g60-touchscreen'],model='G60',country='US',powertrain='BEV',source_label='Owner forum')
    result=import_web(session,request,client=web_client_factory({'https://owners.example/forum/g60-touchscreen':(html,'text/html')}))
    seed_requirements(session);session.commit()
    assert result['imported'][0]['records']>=2
    source=session.scalar(select(Source).where(Source.format=='html',Source.edition=='https://owners.example/forum/g60-touchscreen'))
    assert source and source.origin=='observed' and source.active
    assert Path(source.path).exists() and source.checksum
    refs=session.scalars(select(EvidenceRef).where(EvidenceRef.source_id==source.id)).all()
    assert len(refs)>=2 and refs[0].locator['extraction_method']=='static-html-text-block-v1'
    record=session.scalar(select(Feedback).where(Feedback.source_id==source.id))
    assert record.origin=='observed' and record.source_label=='Owner forum'
    assert session.scalar(select(Aspect).join(Feedback,Aspect.feedback_id==Feedback.id).where(Feedback.source_id==source.id,Aspect.mapping_status=='suggested'))
    assert ranking(session,Cohort()).denominator==3610
    observed=ranking(session,Cohort(origin='observed',powertrain='BEV'))
    assert observed.denominator==source.record_count and any(r.numerator for r in observed.requirements)
    matched=session.execute(text('SELECT evidence_id FROM evidence_fts WHERE evidence_fts MATCH :query'),{'query':'touchscreen'}).scalars().all()
    assert any(ref.id in matched for ref in refs)

def test_web_connector_reimport_versions_and_failures(session,monkeypatch):
    monkeypatch.setattr('backend.connectors.socket.getaddrinfo',lambda host,port:[(socket.AF_INET,socket.SOCK_STREAM,6,'',( '93.184.216.34',0))])
    url='https://owners.example/forum/g60-range'
    first='<p>The charging route planner is difficult to use and I want clearer battery preconditioning status during long trips.</p>'
    second='<p>The charging route planner is still difficult to use, and the battery preconditioning status should be much easier to find before fast charging.</p>'
    request=WebConnectorRequest(urls=[url],model='G60',country='US')
    import_web(session,request,client=web_client_factory({url:(first,'text/html')}));session.commit()
    old=session.scalar(select(Source).where(Source.edition==url,Source.active==True))
    skipped=import_web(session,request,client=web_client_factory({url:(first,'text/html')}));session.commit()
    assert skipped['skipped'] and session.get(Source,old.id).active
    import_web(session,request,client=web_client_factory({url:(second,'text/html')}));session.commit()
    assert session.get(Source,old.id).active is False
    assert session.scalar(select(func.count()).select_from(Source).where(Source.edition==url,Source.active==True))==1
    blocked=web_client_factory({'robots':'User-agent: *\nDisallow: /\n',url:(second,'text/html')})
    with pytest.raises(ValueError,match='did not import any pages'):
        import_web(session,request,client=blocked)
    with pytest.raises(ValueError):
        normalize_url('http://127.0.0.1/private')
    bad=import_web(session,WebConnectorRequest(urls=[url,'https://owners.example/file.pdf']),client=web_client_factory({url:(second,'text/html'),'https://owners.example/file.pdf':('%PDF-1.7','application/pdf')}))
    assert bad['skipped'] and bad['failed'][0]['error'].startswith('Unsupported content type')

def test_pdf_page_preview_is_local_and_exactly_addressed(client,session):
    source=session.scalar(select(Source).where(Source.filename=='G60_OptionList.PDF'))
    response=client.get(f'/api/v1/sources/{source.id}/pages/4')
    assert response.status_code==200 and response.headers['content-type']=='image/png'
    assert response.content.startswith(b'\x89PNG\r\n\x1a\n')
    assert client.get(f'/api/v1/sources/{source.id}/pages/0').status_code==422
    assert client.get(f'/api/v1/sources/{source.id}/pages/999').status_code==404

def test_scenario_requires_reviewed_compatible_bmw_fact(session):
    from backend.comparisons import comparison
    result=comparison(session)
    assert all(r['bmw'] is None and r['comparison_status']=='Not comparable' for r in result['rows'])
    fact=session.scalar(select(Fact).where(Fact.variant=='i5 eDrive40',Fact.label=='System power'))
    fact.review_status='reviewed';session.flush()
    row=next(r for r in comparison(session)['rows'] if r['label']=='System power')
    assert row['bmw']==250 and row['bmw_evidence_id']==fact.evidence_id

def test_deck_empty_approved_and_frozen_snapshot(session):
    from backend.exporter import export_deck
    empty=export_deck(session,ExportRequest(cohort=Cohort()));session.commit()
    record=session.get(Export,empty['export_id']);prs=Presentation(record.path)
    texts=' '.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    assert 'No approved requirements' in texts
    req=target(session);session.rollback()
    decide(session,req.id,decision(req,'approve',band='Medium',rationale='Automated export validation; no real product decision.'))
    result=export_deck(session,ExportRequest(cohort=Cohort()));session.commit()
    saved=session.get(Export,result['export_id']);old_hash=digest(saved.snapshot)
    assert result['approved']==1
    assert all(sample['id'] in saved.snapshot['citations'] for sample in saved.snapshot['approved'][0]['evidence_samples'])
    prs=Presentation(saved.path)
    assert len(prs.slides)==result['slides']
    assert req.title in ' '.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=prs.slide_width+1000
            assert shape.top+shape.height<=prs.slide_height+1000
    session.rollback();decide(session,req.id,decision(req,'edit',title='A subsequent edited requirement',rationale='Verify frozen export remains unchanged.'))
    session.expire_all();assert digest(session.get(Export,result['export_id']).snapshot)==old_hash

def test_long_requirement_tables_paginate_without_truncation():
    from backend.exporter import table_pages,row_height
    widths=[6.1,1.5,3.2,1.2]
    rows=[['A fully preserved long customer requirement title with additional context '*2,'Medium','128/3610 (3.5%)',i] for i in range(18)]
    pages=list(table_pages(rows,widths))
    assert len(pages)>1 and [r for page in pages for r in page]==rows
    assert all(.58+sum(row_height(row,widths) for row in page)<=3.8 for page in pages)

def test_migrations_bootstrap_fresh_database(tmp_path):
    env=os.environ.copy();env['BMWPM_DATA_DIR']=str(tmp_path);env['BMWPM_DATABASE_URL']='sqlite:///'+(tmp_path/'fresh.sqlite').as_posix()
    result=subprocess.run([sys.executable,'-m','backend.cli','migrate'],cwd=ROOT,env=env,capture_output=True,text=True)
    assert result.returncode==0,result.stderr
    with sqlite3.connect(tmp_path/'fresh.sqlite') as db:
        assert db.execute('select version_num from alembic_version').fetchone()[0]=='004_document_rag'
        for table in ('rag_chunks', 'rag_fts', 'rag_runs'):
            assert db.execute("select name from sqlite_master where type='table' and name=?", (table,)).fetchone()
        assert db.execute("select count(*) from sqlite_master where type='trigger'").fetchone()[0]==14
