"""Editable, paginated executive decks from a frozen evidence snapshot."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from .models import Export, Source, EvidenceRef, uid
from .services import ranking, save_priority_run
from .comparisons import comparison
from .config import DATA_DIR
from .audit import append_event, digest
from textwrap import wrap

NAVY='13253D'; BLUE='1769DE'; GRAY='596879'; WHITE='FFFFFF'

def text_box(slide,x,y,w,h,text,size=20,color=NAVY,bold=False):
    shape=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame=shape.text_frame; frame.word_wrap=True
    frame.margin_left=frame.margin_right=0
    frame.margin_top=frame.margin_bottom=0
    for i,line in enumerate(str(text).split('\n')):
        p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
        p.text=line; p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(10)
        p.line_spacing=1.12
    return shape

def slide_base(prs,title,subtitle=None):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    text_box(slide,.65,.45,12,0.35,'BMW PRODUCT INTELLIGENCE',12,BLUE,True)
    text_box(slide,.65,1.02,12,1,title,32,NAVY,True)
    if subtitle: text_box(slide,.65,2.00,12,.65,subtitle,17,GRAY)
    text_box(slide,.65,7.03,11,.25,'Supplied hackathon evidence. Origin unverified. Dates unavailable for feedback.',10,GRAY)
    text_box(slide,12.1,7.03,.5,.25,str(len(prs.slides)),10,GRAY)
    return slide

def row_height(row,widths):
    return max(.58,max(len(wrap(str(value),width=max(8,int(widths[j]*7.5)))) for j,value in enumerate(row))*.28+.18)

def table_pages(rows,widths):
    page=[];height=.58
    for row in rows:
        extra=row_height(row,widths)
        if page and height+extra>3.8:
            yield page;page=[];height=.58
        page.append(row);height+=extra
    if page:yield page

def table(slide,headers,rows,y=2.85,widths=None):
    widths=widths or [12/len(headers)]*len(headers)
    heights=[row_height(row,widths) for row in [headers,*rows]]
    height=sum(heights)
    shape=slide.shapes.add_table(len(rows)+1,len(headers),Inches(.65),Inches(y), Inches(12),Inches(height))
    t=shape.table
    if widths:
        for i,width in enumerate(widths): t.columns[i].width=Inches(width)
    for i,row in enumerate([headers,*rows]):
        t.rows[i].height=Inches(heights[i])
        for j,value in enumerate(row):
            c=t.cell(i,j); c.text=str(value); c.margin_left=Inches(.12); c.margin_top=Inches(.08)
            c.fill.solid(); c.fill.fore_color.rgb=RGBColor.from_string(NAVY if i==0 else 'F3F6FA' if i%2 else WHITE)
            for p in c.text_frame.paragraphs:
                p.font.name='Arial'; p.font.size=Pt(17); p.font.bold=i==0;p.font.color.rgb=RGBColor.from_string(WHITE if i==0 else NAVY)
    return t

def build_deck(snapshot, path):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    cohort=snapshot['cohort']; approvals=snapshot['approved']; all_requirements=snapshot['ranking']['requirements']
    context=f'{cohort["model"]} / {cohort["country"]} / {cohort["powertrain"]} powertrains'
    s=slide_base(prs,'Product decisions',context)
    text_box(s,.65,3,11,1.3,f'{len(approvals)} approved requirements',42,BLUE,True)
    text_box(s,.65,4.5,11,1,f'{snapshot["ranking"]["denominator"]:,} distinct feedback records in this cohort\nSnapshot {snapshot["created_at"][:19]} UTC',21,GRAY)
    if not approvals:
        s=slide_base(prs,'No approved requirements',context)
        text_box(s,.65,3,11,2,'Review the evidence and record PM decisions before presenting an approved requirement matrix.\nThe current snapshot preserves this empty state.',24)
    matrix=[[r['title'],r['band'],f"{r['numerator']}/{r['denominator']} ({(r['demand'] or 0)*100:.1f}%)",r['revision']] for r in approvals]
    for page_index,rows in enumerate(table_pages(matrix,[6.1,1.5,3.2,1.2])):
        s=slide_base(prs,'Approved requirement matrix',f'{context} / page {page_index+1}')
        table(s,['Customer need','Impact','Demand','Rev.'],rows,widths=[6.1,1.5,3.2,1.2])
    for r in approvals:
        s=slide_base(prs,'Approved customer need',f"{r['feature']} / revision {r['revision']} / PM impact {r['band']}")
        text_box(s,.65,2.8,12,1.25,r['title'],24,NAVY,True)
        description=r['description']
        # Split long narratives into full-size continuation slides.
        pieces=wrap_chunks(description,220)
        text_box(s,.65,4.2,12,1.25,pieces[0],21)
        text_box(s,.65,5.85,12,.9,f"Demand: {r['numerator']} distinct supporting records / {r['denominator']} eligible records.\nMapping: {r['mapping_status']}. Engineering feasibility and cost remain unknown.",19,GRAY)
        for piece in pieces[1:]:
            s=slide_base(prs,'Requirement detail',f"{r['id']} / continued")
            text_box(s,.65,2.85,12,3.8,piece,22)
        for index,criterion in enumerate(r['acceptance_criteria']):
            for piece in wrap_chunks(criterion,650):
                s=slide_base(prs,'Proposed acceptance criteria',f'{r["id"]} / criterion {index+1} / proposal awaiting validation')
                text_box(s,.65,2.85,12,3.8,piece,22)
        rationale=r.get('decision_rationale','Approval rationale not available.')
        for piece in wrap_chunks(rationale,650):
            s=slide_base(prs,'PM decision rationale',r['id'])
            text_box(s,.65,2.85,12,3.8,piece,22)
        for sample in r['evidence_samples']:
            for quote_part in wrap_chunks(sample['quote'],550):
                s=slide_base(prs,'Supporting customer evidence',f"{r['feature']} / {sample['external_id']} / {sample['type']}")
                text_box(s,.85,2.85,11.5,3.15,quote_part,23)
                citation=snapshot['citations'].get(sample['id'],sample['id'])
                text_box(s,.65,6.2,12,.55,citation,12,GRAY)
                s.notes_slide.notes_text_frame.text=f'Evidence ID: {sample["id"]}\n{citation}\nOrigin: unknown, supplied hackathon data.'
    s=slide_base(prs,'Evidence limitations and next validation')
    text_box(s,.65,2.65,12,3.9,'Feedback has no timestamps. Velocity is unavailable.\nGerman catalog equipment does not establish US availability.\nStudy instruments remain separate. Sample sizes are unknown.\nSales causality, ROI, engineering effort and cost are unscored.\nNext: review feature mappings, verify the affected variant, and validate proposed acceptance criteria.',22)
    if snapshot.get('scenario'):
        scenario=snapshot['scenario']
        s=slide_base(prs,'Appendix: synthetic competitor scenario','DE / BEV sedan assumptions. Mercedes and BYD entries are fictional.')
        table(s,['Feature / unit','BMW baseline','Mercedes scenario','BYD scenario'],[[f"{r['label']} ({r['unit']})",str(r['bmw']) if r['bmw'] is not None else 'Not comparable',str(r['mercedes']),str(r['byd'])] for r in scenario['rows']],widths=[3.6,2.8,2.8,2.8])
        text_box(s,.65,5.8,12,.9,'Price assumptions include 19% German VAT and exclude delivery.\nSynthetic values establish no actual competitor specifications, demand or ROI.',16,GRAY)
        s.notes_slide.notes_text_frame.text='Generator: authored demonstration fixture v1. Input: G60 catalog, July 2026, page 4.\n'+str(scenario)
    lines=[f'{name}: SHA-256 {checksum[:16]}…' for name,checksum in snapshot['source_register'].items()]
    for start in range(0,len(lines),4):
        s=slide_base(prs,'Snapshot and source register',f'Sources {start+1}–{min(start+4,len(lines))} of {len(lines)}')
        text_box(s,.65,2.8,12,3.8,'\n\n'.join(lines[start:start+4]),17,GRAY)
        s.notes_slide.notes_text_frame.text='Full checksums and immutable evidence snapshot accompany the exported deck in the application.\n'+str(snapshot['source_register'])
    prs.save(path)
    return len(prs.slides)

def wrap_chunks(value,limit):
    chunks=[]; current=''
    for word in value.split():
        if current and len(current)+len(word)+1>limit:
            chunks.append(current); current=''
        current+=(' ' if current else '')+word
    return chunks+[current] if current else chunks or ['']

def export_deck(session, request):
    from .models import Decision, now
    from sqlalchemy import select
    result=ranking(session,request.cohort)
    priority=save_priority_run(session,result)
    approved=[r.model_dump() for r in result.requirements if r.status=='approved']
    for r in approved:
        decisions=session.scalars(select(Decision).where(Decision.requirement_id==r['id'],Decision.revision==r['revision']).order_by(Decision.created_at.desc())).all()
        d=next((d for d in decisions if d.snapshot.get('cohort')==request.cohort.model_dump()),None)
        r['decision_rationale']=d.rationale if d else ''
    citations={}
    for r in approved:
        for sample in r['evidence_samples']:
            ref=session.get(EvidenceRef,sample['id']); source=session.get(Source,ref.source_id)
            citations[ref.id]=f"{source.filename} / {ref.locator.get('sheet','')} / row {ref.locator.get('row','')} / {ref.locator.get('external_id','')}"
    sources=session.scalars(select(Source).where(Source.active==True)).all()
    snapshot={'cohort':request.cohort.model_dump(),'created_at':now(),'ranking':result.model_dump(),'approved':approved,'citations':citations,
              'source_register':{s.filename:s.checksum for s in sources},'scenario':comparison(session) if request.include_scenario else None}
    export_id=uid(); path=DATA_DIR/'exports'/f'BMW-{request.cohort.model}-{export_id[:10]}.pptx'
    slides=build_deck(snapshot,path)
    record=Export(id=export_id,priority_run_id=priority.id,snapshot=snapshot,path=str(path));session.add(record); session.flush()
    append_event(session,'Demo PM','deck.exported',export_id,{'priority_run_id':priority.id,'snapshot_hash':digest(snapshot),'approved':len(approved),'slides':slides})
    return {'export_id':export_id,'slides':slides,'approved':len(approved),'download_url':f'/api/v1/exports/{export_id}/download','snapshot_url':f'/api/v1/exports/{export_id}/snapshot'}
